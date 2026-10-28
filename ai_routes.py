"""
Paid AI tools built on a saved comparison: bullet rewrites, cover letter, interview
prep, tailored DOCX export, and bulk ranking for coaches.
"""
import asyncio
import io
import logging
import re
from typing import List, Optional

import docx
from docx.shared import Pt
from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field

from compare_routes import load_comparison
from database import get_comparisons_collection
from deps import get_active_user, require_feature
from file_utils import extract_document
from openai_client import (
    AIServiceError,
    call_openai_completions,
    generate_cover_letter,
    generate_interview_prep,
    generate_rewrites,
)
from plans import get_plan
from usage import AI_ACTIONS, SCANS, consume, enforce_rate_limit, refund

logger = logging.getLogger(__name__)

router = APIRouter(tags=["AI tools"])


async def _run_metered(user: dict, coro_fn, *args):
    """Reserve an AI action, run the call, and give the action back if the AI service failed"""
    await enforce_rate_limit(f"ai:{user['id']}", 20, 60, "You are going too fast. Please wait a minute.")
    await consume(user, AI_ACTIONS)
    try:
        return await coro_fn(*args)
    except AIServiceError as e:
        await refund(user, AI_ACTIONS)
        raise HTTPException(status_code=502, detail=str(e))
    except ValueError as e:
        await refund(user, AI_ACTIONS)
        logger.error(f"AI configuration error: {e}")
        raise HTTPException(status_code=500, detail="The service is not configured correctly. Please contact support.")
    except Exception as e:
        await refund(user, AI_ACTIONS)
        logger.error(f"Unexpected AI error: {e}", exc_info=True)
        raise HTTPException(status_code=502, detail="The AI service is temporarily unavailable. Please try again.")


class ComparisonRef(BaseModel):
    comparison_id: str
    regenerate: bool = False


class CoverLetterRequest(ComparisonRef):
    tone: str = Field("professional", pattern="^(professional|friendly|confident|concise)$")


@router.get("/api/ai/{comparison_id}")
async def get_saved_ai_outputs(comparison_id: str, user: dict = Depends(get_active_user)):
    """Previously generated rewrites / cover letter / interview prep for a comparison"""
    doc = await load_comparison(user, comparison_id)
    return {
        "rewrites": doc.get("rewrites"),
        "cover_letter": doc.get("cover_letter"),
        "interview_prep": doc.get("interview_prep"),
    }


@router.post("/api/ai/rewrites")
async def rewrites(body: ComparisonRef, user: dict = Depends(require_feature("rewrites"))):
    doc = await load_comparison(user, body.comparison_id, with_text=True)
    if doc.get("rewrites") and not body.regenerate:
        return doc["rewrites"]
    result = await _run_metered(user, generate_rewrites, doc["jd_text"], doc["resume_text"], doc.get("missing_skills", []))
    await get_comparisons_collection().update_one({"_id": doc["_id"]}, {"$set": {"rewrites": result}})
    return result


@router.post("/api/ai/cover-letter")
async def cover_letter(body: CoverLetterRequest, user: dict = Depends(require_feature("cover_letter"))):
    doc = await load_comparison(user, body.comparison_id, with_text=True)
    cached = doc.get("cover_letter")
    if cached and cached.get("tone") == body.tone and not body.regenerate:
        return cached
    letter = await _run_metered(
        user, generate_cover_letter, doc["jd_text"], doc["resume_text"], body.tone,
        doc.get("job_title") or "", doc.get("company") or "",
    )
    result = {"letter": letter, "tone": body.tone}
    await get_comparisons_collection().update_one({"_id": doc["_id"]}, {"$set": {"cover_letter": result}})
    return result


@router.post("/api/ai/interview-prep")
async def interview_prep(body: ComparisonRef, user: dict = Depends(require_feature("interview_prep"))):
    doc = await load_comparison(user, body.comparison_id, with_text=True)
    if doc.get("interview_prep") and not body.regenerate:
        return {"questions": doc["interview_prep"]}
    questions = await _run_metered(user, generate_interview_prep, doc["jd_text"], doc["resume_text"], doc.get("missing_skills", []))
    await get_comparisons_collection().update_one({"_id": doc["_id"]}, {"$set": {"interview_prep": questions}})
    return {"questions": questions}


# ============ Tailored DOCX export ============

class AcceptedRewrite(BaseModel):
    original: str = Field(..., max_length=600)
    suggested: str = Field(..., max_length=600)


class ExportRequest(BaseModel):
    comparison_id: str
    accepted: List[AcceptedRewrite] = Field(default_factory=list, max_length=20)


HEADING_RE = re.compile(
    r"^(summary|profile|objective|experience|work experience|professional experience|employment|education|"
    r"skills|technical skills|projects|certifications|awards|publications|languages|interests)\s*:?$", re.IGNORECASE
)
BULLET_PREFIX_RE = re.compile(r"^\s*([-*•▪●◦‣·])\s+")


def apply_rewrites(text: str, accepted: List[AcceptedRewrite]) -> str:
    """Swap accepted original lines for their rewrites (whole-line match first, then substring)"""
    for item in accepted:
        original, suggested = item.original.strip(), item.suggested.strip()
        if not original or not suggested:
            continue
        lines = text.split("\n")
        replaced = False
        for i, line in enumerate(lines):
            bare = BULLET_PREFIX_RE.sub("", line).strip()
            if bare == original:
                prefix = BULLET_PREFIX_RE.match(line)
                lines[i] = (prefix.group(0) if prefix else "") + suggested
                replaced = True
                break
        text = "\n".join(lines) if replaced else text.replace(original, suggested, 1)
    return text


def build_docx(text: str) -> bytes:
    """Rebuild the resume as a clean single-column Word document that ATS parsers read well"""
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for index, line in enumerate(text.split("\n")):
        stripped = line.strip()
        if not stripped:
            continue
        if BULLET_PREFIX_RE.match(line):
            document.add_paragraph(BULLET_PREFIX_RE.sub("", line).strip(), style="List Bullet")
        elif HEADING_RE.match(stripped) or (stripped.isupper() and len(stripped) < 40):
            document.add_heading(stripped.title() if stripped.isupper() else stripped, level=2)
        elif index == 0:
            document.add_heading(stripped, level=1)
        else:
            document.add_paragraph(stripped)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@router.post("/api/export/tailored")
async def export_tailored(body: ExportRequest, user: dict = Depends(require_feature("export"))):
    doc = await load_comparison(user, body.comparison_id, with_text=True)
    tailored = apply_rewrites(doc["resume_text"], body.accepted)
    content = build_docx(tailored)
    filename = re.sub(r"[^A-Za-z0-9_-]+", "_", f"resume_{doc.get('company') or doc.get('job_title') or 'tailored'}")[:60] + ".docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ============ Bulk ranking (Coach plan) ============

@router.post("/api/bulk/rank")
async def bulk_rank(
    jd_text: str = Form(...),
    resume_files: List[UploadFile] = File(...),
    user: dict = Depends(require_feature("bulk_rank")),
):
    """Rank several resumes against one job description. Nothing is stored."""
    plan = get_plan(user)
    if len(jd_text.strip()) < 10:
        raise HTTPException(status_code=400, detail="The job description is too short")
    if not resume_files or len(resume_files) > plan.bulk_max_resumes:
        raise HTTPException(status_code=400, detail=f"Upload between 1 and {plan.bulk_max_resumes} resumes")

    await enforce_rate_limit(f"bulk:{user['id']}", 6, 60, "Too many bulk runs. Please wait a minute.")
    await consume(user, SCANS, len(resume_files))

    semaphore = asyncio.Semaphore(5)

    async def score(upload: UploadFile) -> dict:
        name = upload.filename or "resume"
        try:
            doc = await extract_document(upload)
        except ValueError as e:
            return {"name": name, "error": str(e), "refund": True}
        except Exception:
            return {"name": name, "error": "Could not read this file", "refund": True}
        async with semaphore:
            try:
                result = await call_openai_completions(jd_text.strip(), doc.text)
            except Exception as e:
                logger.error(f"Bulk scoring failed for {name}: {e}")
                return {"name": name, "error": "Scoring failed, try again", "refund": True}
        return {
            "name": name,
            "matchPercent": result.matchPercent,
            "matchedSkills": result.matchedSkills,
            "missingSkills": result.missingSkills,
        }

    results = await asyncio.gather(*(score(f) for f in resume_files))
    failed = sum(1 for r in results if r.pop("refund", False))
    if failed:
        await refund(user, SCANS, failed)
    ranked = sorted((r for r in results if "matchPercent" in r), key=lambda r: r["matchPercent"], reverse=True)
    for position, item in enumerate(ranked, start=1):
        item["rank"] = position
    return {"ranked": ranked, "errors": [r for r in results if "error" in r]}
