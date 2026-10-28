"""POST /api/compare, history, saved resumes and plan shaping"""
import io

import docx
from unittest.mock import AsyncMock

from openai_client import AIServiceError
from tests.conftest import fake_compare_result

JD = "We need a Backend Engineer with Python, Docker and AWS experience for our platform team."
RESUME = (b"Jane Doe\njane@example.com +1 555 123 4567\nExperience\n- Built Python services 2019-2023\n"
          b"- Used Docker and Postgres 2020-2022\nEducation\nBSc Computer Science 2015\nSkills\nPython Docker\n" * 2)


def resume_files(content=RESUME, name="resume.txt", ctype="text/plain"):
    return {"resume_file": (name, content, ctype)}


async def test_requires_auth(client):
    r = await client.post("/api/compare", data={"jd_text": JD}, files=resume_files())
    assert r.status_code in (401, 403)


async def test_free_compare_hides_paid_data(client, make_user, mock_ai):
    headers, _ = await make_user()
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["matchPercent"] == 72 and body["id"]
    assert body["breakdown"] is None
    assert set(body["locked"]) == {"breakdown", "ats_details"}
    assert body["ats"]["locked"] is True and body["ats"]["issues"] == []
    assert body["usage"] == {"used": 1, "limit": 3}


async def test_pro_compare_includes_breakdown_and_ats_details(client, make_user, mock_ai):
    headers, _ = await make_user("pro")
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    body = r.json()
    assert body["breakdown"]["titleMatch"] == 90
    assert body["locked"] == []
    assert body["ats"]["locked"] is False and "passed" in body["ats"]


async def test_free_quota_then_402(client, make_user, mock_ai):
    headers, _ = await make_user()
    for _ in range(3):
        assert (await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())).status_code == 200
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    assert r.status_code == 402
    assert r.json()["detail"]["code"] == "quota_exceeded"


async def test_failed_ai_call_does_not_use_a_scan(client, make_user, mock_ai):
    headers, uid = await make_user()
    mock_ai["compare"].side_effect = AIServiceError("The AI service is temporarily unavailable. Please try again.")
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    assert r.status_code == 502
    assert "temporarily unavailable" in r.json()["detail"]
    mock_ai["compare"].side_effect = None
    status = (await client.get("/api/billing/status", headers=headers)).json()
    assert status["usage"]["scans"] == 0


async def test_internal_errors_are_not_leaked(client, make_user, mock_ai):
    headers, _ = await make_user()
    mock_ai["compare"].side_effect = RuntimeError("secret internal detail sk-12345")
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    assert r.status_code == 502
    assert "sk-12345" not in r.text


async def test_invalid_inputs(client, make_user, mock_ai):
    headers, _ = await make_user()
    no_jd = await client.post("/api/compare", headers=headers, files=resume_files())
    assert no_jd.status_code == 400
    no_resume = await client.post("/api/compare", headers=headers, data={"jd_text": JD})
    assert no_resume.status_code == 400
    short = await client.post("/api/compare", headers=headers, data={"jd_text": "short"}, files=resume_files())
    assert short.status_code == 400
    bad_type = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files(b"x", "r.png", "image/png"))
    assert bad_type.status_code == 400
    legacy = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files(b"x" * 100, "r.doc", "application/msword"))
    assert legacy.status_code == 400 and ".docx" in legacy.json()["detail"]
    # None of the rejected requests used a scan
    assert (await client.get("/api/billing/status", headers=headers)).json()["usage"]["scans"] == 0


async def test_file_too_large(client, make_user, mock_ai, monkeypatch):
    import file_utils
    monkeypatch.setattr(file_utils.settings, "MAX_FILE_SIZE", 50)
    headers, _ = await make_user()
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files(b"a" * 500))
    assert r.status_code == 400 and "exceeds" in r.json()["detail"]


async def test_unusable_model_output_is_not_charged(client, make_user, mock_ai):
    headers, _ = await make_user()
    mock_ai["compare"].return_value = fake_compare_result(matchPercent=0, matchedSkills=[], missingSkills=[], breakdown=None, warnings=["The AI response could not be read. Please try again."])
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())
    assert r.status_code == 200 and r.json()["usage"]["used"] == 0


async def test_docx_resume_gets_structural_ats_findings(client, make_user, mock_ai):
    d = docx.Document()
    d.add_paragraph("Jane Doe jane@example.com")
    d.add_paragraph("Experience 2019-2023 Built Python services")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Skills"
    t.cell(0, 1).text = "Python Docker AWS"
    buf = io.BytesIO()
    d.save(buf)
    headers, _ = await make_user("pro")
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD},
                          files=resume_files(buf.getvalue(), "cv.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    titles = [i["title"] for i in r.json()["ats"]["issues"]]
    assert "Tables detected" in titles


async def test_history_list_detail_and_free_limit(client, make_user, mock_ai):
    headers, _ = await make_user("pro")
    ids = []
    for _ in range(7):
        ids.append((await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files=resume_files())).json()["id"])
    listing = (await client.get("/api/history", headers=headers)).json()
    assert listing["total"] == 7 and listing["items"][0]["id"] == ids[-1]
    assert not any(i["locked"] for i in listing["items"])
    assert (await client.get(f"/api/history/{ids[0]}", headers=headers)).status_code == 200

    # Same data viewed on the free plan: only the latest 5 open
    from database import get_users_collection
    uid = (await client.get("/api/auth/me", headers=headers)).json()["id"]
    await get_users_collection().update_one({"_id": uid}, {"$set": {"subscription_status": "canceled"}})
    listing = (await client.get("/api/history", headers=headers)).json()
    assert [i["locked"] for i in listing["items"]] == [False] * 5 + [True] * 2
    assert (await client.get(f"/api/history/{ids[0]}", headers=headers)).status_code == 402
    assert (await client.get(f"/api/history/{ids[-1]}", headers=headers)).status_code == 200


async def test_history_is_private_to_the_owner(client, make_user, mock_ai):
    h1, _ = await make_user()
    h2, _ = await make_user()
    cid = (await client.post("/api/compare", headers=h1, data={"jd_text": JD}, files=resume_files())).json()["id"]
    assert (await client.get(f"/api/history/{cid}", headers=h2)).status_code == 404
    assert (await client.request("DELETE", f"/api/history/{cid}", headers=h2)).status_code == 404
    assert (await client.request("DELETE", f"/api/history/{cid}", headers=h1)).status_code == 200


async def test_saved_resumes_limit_and_reuse(client, make_user, mock_ai):
    headers, _ = await make_user()
    up = await client.post("/api/resumes", headers=headers, data={"name": "Main"}, files=resume_files())
    assert up.status_code == 201
    rid = up.json()["id"]
    assert up.json()["words"] > 10
    second = await client.post("/api/resumes", headers=headers, files=resume_files())
    assert second.status_code == 402 and second.json()["detail"]["code"] == "resume_limit"

    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD, "resume_id": rid})
    assert r.status_code == 200
    assert len((await client.get("/api/resumes", headers=headers)).json()["items"]) == 1
    assert (await client.request("DELETE", f"/api/resumes/{rid}", headers=headers)).status_code == 200
    gone = await client.post("/api/compare", headers=headers, data={"jd_text": JD, "resume_id": rid})
    assert gone.status_code == 404


async def test_jd_url_requires_paid_plan(client, make_user, mock_ai):
    headers, _ = await make_user()
    r = await client.post("/api/compare", headers=headers, data={"jd_url": "https://example.com/job"}, files=resume_files())
    assert r.status_code == 402
    assert (await client.post("/api/import-job", headers=headers, json={"url": "https://example.com/job"})).status_code == 402


async def test_pro_jd_url_import_feeds_compare(client, make_user, mock_ai, monkeypatch):
    monkeypatch.setattr("compare_routes.fetch_job_text", AsyncMock(return_value={"text": JD, "title": "Backend Engineer", "company": "Acme", "url": "https://x.io/j"}))
    headers, _ = await make_user("pro")
    r = await client.post("/api/compare", headers=headers, data={"jd_url": "https://x.io/j"}, files=resume_files())
    assert r.status_code == 200
    sent_jd = mock_ai["compare"].call_args.args[0]
    assert "Backend Engineer" in sent_jd
