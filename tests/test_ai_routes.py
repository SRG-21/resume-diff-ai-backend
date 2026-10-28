"""Rewrites, cover letter, interview prep, export, bulk ranking and the tracker"""
import io

import docx

from tests.conftest import fake_compare_result
from ai_routes import AcceptedRewrite, apply_rewrites, build_docx

JD = "We need a Backend Engineer with Python, Docker and AWS experience for our platform team."
RESUME = (b"Jane Doe\njane@example.com\nExperience\n- Worked on Python apps\n- Used Docker\nEducation\nBSc 2015\n")


async def new_comparison(client, headers):
    r = await client.post("/api/compare", headers=headers, data={"jd_text": JD}, files={"resume_file": ("r.txt", RESUME, "text/plain")})
    assert r.status_code == 200, r.text
    return r.json()["id"]


async def test_ai_tools_are_paid_only(client, make_user, mock_ai):
    headers, _ = await make_user()
    cid = await new_comparison(client, headers)
    for path, body in [
        ("/api/ai/rewrites", {"comparison_id": cid}),
        ("/api/ai/cover-letter", {"comparison_id": cid}),
        ("/api/ai/interview-prep", {"comparison_id": cid}),
        ("/api/export/tailored", {"comparison_id": cid}),
    ]:
        r = await client.post(path, headers=headers, json=body)
        assert r.status_code == 402, path
        assert r.json()["detail"]["code"] == "upgrade_required"
    mock_ai["rewrites"].assert_not_awaited()


async def test_rewrites_cached_and_metered(client, make_user, mock_ai):
    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    r1 = await client.post("/api/ai/rewrites", headers=headers, json={"comparison_id": cid})
    assert r1.status_code == 200 and r1.json()["rewrites"][0]["suggested"].startswith("Built")
    r2 = await client.post("/api/ai/rewrites", headers=headers, json={"comparison_id": cid})
    assert r2.json() == r1.json()
    assert mock_ai["rewrites"].await_count == 1  # second call served from cache
    assert (await client.get("/api/billing/status", headers=headers)).json()["usage"]["ai_actions"] == 1
    await client.post("/api/ai/rewrites", headers=headers, json={"comparison_id": cid, "regenerate": True})
    assert mock_ai["rewrites"].await_count == 2

    saved = (await client.get(f"/api/ai/{cid}", headers=headers)).json()
    assert saved["rewrites"] and saved["cover_letter"] is None


async def test_ai_failure_refunds_the_action(client, make_user, mock_ai):
    from openai_client import AIServiceError
    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    mock_ai["cover"].side_effect = AIServiceError("The AI service is temporarily unavailable. Please try again.")
    r = await client.post("/api/ai/cover-letter", headers=headers, json={"comparison_id": cid})
    assert r.status_code == 502
    assert (await client.get("/api/billing/status", headers=headers)).json()["usage"]["ai_actions"] == 0


async def test_cover_letter_and_interview_prep(client, make_user, mock_ai):
    headers, _ = await make_user("pass")
    cid = await new_comparison(client, headers)
    letter = await client.post("/api/ai/cover-letter", headers=headers, json={"comparison_id": cid, "tone": "friendly"})
    assert letter.json() == {"letter": "Dear [Hiring Manager], ...", "tone": "friendly"}
    assert mock_ai["cover"].call_args.args[2] == "friendly"
    bad_tone = await client.post("/api/ai/cover-letter", headers=headers, json={"comparison_id": cid, "tone": "rude"})
    assert bad_tone.status_code == 422
    prep = await client.post("/api/ai/interview-prep", headers=headers, json={"comparison_id": cid})
    assert prep.json()["questions"][0]["question"].startswith("Tell me")


async def test_cannot_use_another_users_comparison(client, make_user, mock_ai):
    h1, _ = await make_user("pro")
    h2, _ = await make_user("pro")
    cid = await new_comparison(client, h1)
    assert (await client.post("/api/ai/rewrites", headers=h2, json={"comparison_id": cid})).status_code == 404
    assert (await client.get(f"/api/ai/{cid}", headers=h2)).status_code == 404


def test_apply_rewrites_replaces_whole_bullet_lines():
    text = "Experience\n- Worked on Python apps\n- Used Docker"
    out = apply_rewrites(text, [AcceptedRewrite(original="Worked on Python apps", suggested="Built Python services")])
    assert out == "Experience\n- Built Python services\n- Used Docker"


def test_apply_rewrites_substring_fallback_and_missing_original():
    out = apply_rewrites("I used Docker daily", [AcceptedRewrite(original="used Docker", suggested="ran Docker"), AcceptedRewrite(original="nonexistent", suggested="x")])
    assert out == "I ran Docker daily"


def test_build_docx_structure():
    data = build_docx("Jane Doe\nEXPERIENCE\n- Built things\nPlain paragraph")
    d = docx.Document(io.BytesIO(data))
    styles = [(p.style.name, p.text) for p in d.paragraphs]
    assert ("Heading 1", "Jane Doe") in styles
    assert ("Heading 2", "Experience") in styles
    assert ("List Bullet", "Built things") in styles
    assert ("Normal", "Plain paragraph") in styles


async def test_export_tailored_docx(client, make_user, mock_ai):
    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    r = await client.post("/api/export/tailored", headers=headers, json={
        "comparison_id": cid,
        "accepted": [{"original": "Worked on Python apps", "suggested": "Built Python services handling 1M requests/day"}],
    })
    assert r.status_code == 200
    assert r.headers["content-disposition"].endswith('.docx"')
    d = docx.Document(io.BytesIO(r.content))
    texts = [p.text for p in d.paragraphs]
    assert "Built Python services handling 1M requests/day" in texts
    assert "Worked on Python apps" not in texts


async def test_bulk_rank_coach_only(client, make_user, mock_ai):
    pro_headers, _ = await make_user("pro")
    files = [("resume_files", ("a.txt", RESUME, "text/plain"))]
    assert (await client.post("/api/bulk/rank", headers=pro_headers, data={"jd_text": JD}, files=files)).status_code == 402

    headers, _ = await make_user("coach")
    scores = iter([40, 90, 65])

    async def fake(jd, resume):
        return fake_compare_result(matchPercent=next(scores))

    mock_ai["compare"].side_effect = fake
    files = [("resume_files", (n, RESUME, "text/plain")) for n in ("a.txt", "b.txt", "c.txt")] + [("resume_files", ("bad.png", b"x", "image/png"))]
    r = await client.post("/api/bulk/rank", headers=headers, data={"jd_text": JD}, files=files)
    body = r.json()
    assert [x["matchPercent"] for x in body["ranked"]] == [90, 65, 40]
    assert [x["rank"] for x in body["ranked"]] == [1, 2, 3]
    assert body["errors"][0]["name"] == "bad.png"
    # 3 scans charged, the unreadable file was refunded
    assert (await client.get("/api/billing/status", headers=headers)).json()["usage"]["scans"] == 3


async def test_bulk_rank_enforces_file_cap(client, make_user, mock_ai):
    headers, _ = await make_user("coach")
    files = [("resume_files", (f"{i}.txt", RESUME, "text/plain")) for i in range(11)]
    assert (await client.post("/api/bulk/rank", headers=headers, data={"jd_text": JD}, files=files)).status_code == 400


async def test_tracker_crud_and_gating(client, make_user, mock_ai):
    free, _ = await make_user()
    assert (await client.get("/api/applications", headers=free)).status_code == 402

    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    created = await client.post("/api/applications", headers=headers, json={"company": "Acme", "role": "Backend Engineer", "comparison_id": cid})
    assert created.status_code == 201 and created.json()["match_percent"] == 72 and created.json()["status"] == "saved"
    app_id = created.json()["id"]

    bad = await client.post("/api/applications", headers=headers, json={"company": "X", "role": "Y", "status": "weird"})
    assert bad.status_code == 422
    missing_cmp = await client.post("/api/applications", headers=headers, json={"company": "X", "role": "Y", "comparison_id": "nope"})
    assert missing_cmp.status_code == 404

    upd = await client.patch(f"/api/applications/{app_id}", headers=headers, json={"status": "interview", "notes": "Phone screen Tue"})
    assert upd.json()["status"] == "interview" and upd.json()["notes"] == "Phone screen Tue"
    assert (await client.get("/api/applications", headers=headers)).json()["items"][0]["status"] == "interview"

    other, _ = await make_user("pro")
    assert (await client.patch(f"/api/applications/{app_id}", headers=other, json={"status": "offer"})).status_code == 404
    assert (await client.request("DELETE", f"/api/applications/{app_id}", headers=headers)).status_code == 200
    assert (await client.get("/api/applications", headers=headers)).json()["items"] == []


async def test_export_tailored_pdf_is_text_based_and_applies_rewrites(client, make_user, mock_ai):
    import pypdf

    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    r = await client.post("/api/export/tailored", headers=headers, json={
        "comparison_id": cid, "format": "pdf",
        "accepted": [{"original": "Worked on Python apps", "suggested": "Built Python \u201cservices\u201d \u2014 1M req/day \u4e2d"}],
    })
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf" and r.headers["content-disposition"].endswith('.pdf"')
    text = "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(r.content)).pages)
    assert 'Built Python "services" - 1M req/day' in text  # smart quotes/dash mapped, CJK replaced rather than crashing
    assert "Worked on Python apps" not in text
    assert "jane@example.com" in text


async def test_export_rejects_unknown_format(client, make_user, mock_ai):
    headers, _ = await make_user("pro")
    cid = await new_comparison(client, headers)
    r = await client.post("/api/export/tailored", headers=headers, json={"comparison_id": cid, "format": "rtf"})
    assert r.status_code == 422


def test_build_pdf_handles_escaping_and_empty_input():
    from ai_routes import build_pdf
    assert build_pdf("A & B <script>\n- item <b>x</b>").startswith(b"%PDF")
    assert build_pdf("").startswith(b"%PDF")
